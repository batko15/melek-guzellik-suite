#!/usr/bin/env python3
"""Alien 1 - Datenanalyse: Extrahiert Geschäftdaten aus den AutoFaszination-Uploads."""
import json, os

SRC = "/home/z/my-project/upload/autofaszination-extracted"
OUT = "/home/z/my-project/scripts/extracted"
os.makedirs(OUT, exist_ok=True)

# ---------- 1. Markenhäuser Excel ----------
import openpyxl
xlsx_path = os.path.join(SRC, "Markenhaeuser_vollstaendig_sortiert(2)(1).xlsx")
wb = openpyxl.load_workbook(xlsx_path, data_only=True)
print("=== Excel Sheets:", wb.sheetnames)
for sheet_name in wb.sheetnames[:3]:
    ws = wb[sheet_name]
    print(f"\n--- Sheet '{sheet_name}': {ws.max_row} Zeilen x {ws.max_column} Spalten")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 8: break
        print([str(c)[:30] if c is not None else "" for c in row])
